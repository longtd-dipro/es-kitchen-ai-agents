# -*- coding: utf-8 -*-
"""
メッセージ・メールのまとめ（すべての画面設計書をまたぐ）を作る。make.py から呼ぶ。
  - データ/_共通_メッセージ.py（共通の元データ）と、各画面のデータの使い方を集める
  - 文言は2つ：ja_c＝法人・拠点向け（丁寧）、ja_o＝社内向け（運営・ドライバー・委託・仕入先・倉庫）
  - Message List（Excel・正）を読んで、登録済み／文言の違い／未登録 を自動で判定する
  - 出力/_共通/メッセージ・メール一覧.xlsx と .html を作る
"""
import glob, json, os, re

HERE = os.path.dirname(os.path.abspath(__file__))
TYPE_ML = {'エラー': 'エラー (error)', '警告': '警告 (warning)', '情報': '情報 (Info)', '成功': '成功 (Success)', '確認': '情報 (Info)'}

def ml_format(where):
    w = where or ''
    if 'モーダル' in w: return 'モーダルウィンドウ'
    if 'トースト' in w: return 'トーストメッセージ'
    if 'インライン' in w or '項目の下' in w or '表の下' in w or '後ろ' in w: return 'インラインメッセージ'
    if 'お知らせ' in w: return 'お知らせ一覧'
    return 'ページ内メッセージエリア'

def norm(t):
    t = re.sub(r'\{[^}]*\}', '{}', str(t or ''))
    return re.sub(r'[\s　]+', '', t).replace('。', '').replace('／', '/')

# ---------------------------------------------------------------- Message List を読む
# 正は Excel。cloud など Excel が見えない所では、ml_export.py で書き出した CSV の写し（MessageList_WEB_<日付>.csv・MessageList_MAIL_<日付>.csv）を読む
CSV_WEB, CSV_MAIL = 'MessageList_WEB_*.csv', 'MessageList_MAIL_*.csv'

def parse_web(rows):
    """「運営管理者／法人顧客向けWEB」シートの行（1行目＝見出し）→ {No: {...}}"""
    out = {}
    for r, row in enumerate(rows, start=1):
        if r < 2: continue
        if not row or len(row) < 7 or not row[6]: continue
        n = toint(row[0]) or r - 1
        out[n] = {'phase': row[1], 'scene': row[2], 'screen': row[3], 'button': row[4], 'title': row[5],
                  'ja': str(row[6]), 'type': row[8] if len(row) > 8 else '', 'fmt': row[9] if len(row) > 9 else ''}
    return out

def parse_mail(rows):
    """「メール」シートの行（1〜2行目＝見出し）→ {No: {...}}"""
    out = {}
    for r, row in enumerate(rows, start=1):
        if r < 3: continue
        n = toint(row[0]) if row else None
        if n is None: continue
        out[n] = {'use': str(row[1] or ''), 'to': str(row[3] or ''), 'subject': str(row[4] or '')}
    return out

def read_xlsx_sheets(path):
    """Excel の WEB／メールのシートを行のまま返す {'web': rows, 'mail': rows}。式ではなく値（Google 翻訳の式などを CSV に入れない）"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        if 'WEB' in ws.title and 'web' not in sheets:
            sheets['web'] = [list(r) for r in ws.iter_rows(values_only=True)]
        elif 'メール' in ws.title and 'mail' not in sheets:
            sheets['mail'] = [list(r) for r in ws.iter_rows(values_only=True)]
    return sheets

def read_csv_rows(path):
    import csv
    with open(path, encoding='utf-8-sig', newline='') as f:
        return [[c if c != '' else None for c in row] for row in csv.reader(f)]

def read_ml(root, pattern):
    """戻り値 {'file':…, 'web': {n: {...}}, 'mail': {n: {...}}}。見つからなければ None
    Excel（pattern）があればそれを、なければ同じフォルダの CSV の写し（最新の日付）を読む"""
    if not (root and pattern):
        return None
    files = sorted(glob.glob(os.path.join(root, pattern)))
    if files:
        path = files[-1]
        sh = read_xlsx_sheets(path)
        return {'file': os.path.relpath(path, root), 'web': parse_web(sh.get('web', [])), 'mail': parse_mail(sh.get('mail', []))}
    d = os.path.join(root, os.path.dirname(pattern))
    web = sorted(glob.glob(os.path.join(d, CSV_WEB))); mail = sorted(glob.glob(os.path.join(d, CSV_MAIL)))
    if not web:
        return None
    out = {'file': os.path.relpath(web[-1], root) + '（CSV の写し）', 'web': parse_web(read_csv_rows(web[-1])), 'mail': {}}
    if mail: out['mail'] = parse_mail(read_csv_rows(mail[-1]))
    return out

def toint(v):
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return None

def ml_no(ml, kind):
    m = re.match(r'%s-(\d+)$' % kind, ml or '')
    return int(m.group(1)) if m else None

def both(m):
    """一覧に出す文言（法人向けと社内向けが違えば2行）"""
    c, o = m.get('ja_c', ''), m.get('ja_o', '')
    if c and o and c != o: return '【法人・拠点】%s\n【社内】%s' % (c, o)
    return c or o or m.get('ja', '')

def enrich(common, ml):
    """メッセージに Message List の文言と状態を付ける（各画面の出力にも使う）"""
    for k, m in common.MESSAGES.items():
        m['ja'] = both(m)
        n = ml_no(m.get('ml'), 'WEB')
        m['mlText'], m['mlState'] = '', '未登録（追加が必要）'
        if n is not None:
            hit = ml and ml['web'].get(n)
            if not ml:
                m['mlState'] = '登録済み（Message List を読めなかったので文言は未確認）'
            elif not hit:
                m['mlState'] = 'Message List に %s がない' % m['ml']
            else:
                m['mlText'] = hit['ja']
                m['mlState'] = '登録済み（一致）' if norm(hit['ja']) in [norm(m.get(x)) for x in ('ja_c', 'ja_o') if m.get(x)] else '登録済み（文言が違う）'
    for k, m in common.MAILS.items():
        n = ml_no(m.get('ml'), 'MAIL')
        m['mlSubject'] = (ml['mail'].get(n, {}).get('subject', '') if (ml and n) else '')
        m['mlState'] = '未登録（追加が必要）' if n is None else ('登録済' if m['mlSubject'] or not ml else 'Message List に %s がない' % m['ml'])

# ---------------------------------------------------------------- まとめ
def usage(datas):
    use = {}
    for D in datas:
        names = {x['code']: x['ja'] for x in D.SCREENS}
        for v in D.VIEWS:
            for it in v['items']:
                for c in it.get('err', []):
                    use.setdefault(c, []).append({'doc': D.TITLE[0], 'code': v['code'], 'sname': names.get(v['code'], ''), 'view': v['id'], 'no': it['no'], 'ja': it['ja'], 'site': getattr(D, 'SITE', '')})
    return use

def run(common, datas, ml, out):
    os.makedirs(out, exist_ok=True)
    use = usage(datas)
    msgs, mails, adds, unused = [], [], [], []
    for k, m in common.MESSAGES.items():
        u = use.get(k, [])
        if not u: unused.append(k)
        msgs.append(dict(id=k, use=u, **m))
        if m['mlState'].startswith('未登録'):
            adds.append({'id': k, 'phase': 'フェーズ２', 'scene': m.get('scene', ''),
                         'screen': '、'.join(sorted(set('%s %s' % (x['code'], x['sname']) for x in u))) or '（未使用）',
                         'button': m.get('button', ''), 'title': m['ja'].split('／')[0] if m.get('type') == '確認' else '-',
                         'ja': m['ja'].split('／', 1)[1] if m.get('type') == '確認' and '／' in m['ja'] else m['ja'],
                         'popup': m.get('popup', '-') or '-', 'type': TYPE_ML.get(m.get('type'), m.get('type')), 'fmt': ml_format(m.get('where'))})
    used_mail = set()
    for k, m in common.MAILS.items():
        mails.append(dict(id=k, use=use.get(k, []), **m))
        n = ml_no(m.get('ml'), 'MAIL')
        if n: used_mail.add(n)
    ml_only = [] if not ml else [dict(no=n, **x) for n, x in sorted(ml['mail'].items()) if n not in used_mail]
    data = {'mlFile': (ml or {}).get('file', ''), 'messages': msgs, 'mails': mails, 'adds': adds, 'mlOnly': ml_only, 'unused': unused,
            'docs': [{'title': D.TITLE, 'base': D.BASENAME, 'out': getattr(D, 'OUT_DIR', D.BASENAME)} for D in datas]}
    build_excel(data, os.path.join(out, 'メッセージ・メール一覧.xlsx'))
    build_html(data, os.path.join(out, 'メッセージ・メール一覧.html'))
    return {'messages': len(msgs), 'adds': len(adds), 'diff': sum(1 for m in msgs if '違う' in m['mlState']),
            'mails': len(mails), 'mlOnly': len(ml_only), 'unused': unused, 'mlFile': data['mlFile']}

def where_text(u):
    return '\n'.join('%s %s No.%s %s' % (x['code'], x['view'], x['no'], x['ja']) for x in u)

def build_excel(d, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as L
    wb = Workbook()
    HDR = PatternFill('solid', fgColor='92CDDC'); WARN = PatternFill('solid', fgColor='FFF2CC'); RED = PatternFill('solid', fgColor='FCE4E4'); OK = PatternFill('solid', fgColor='E3F5EE')
    thin = Side(style='thin', color='A6B4C3'); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
    WR = Alignment(wrap_text=True, vertical='top'); CT = Alignment(horizontal='center', vertical='top', wrap_text=True)
    def sheet(ws, cols, rows, fills=None):
        for j, (h, w) in enumerate(cols):
            c = ws.cell(row=1, column=j + 1, value=h); c.font = Font(name='Meiryo UI', bold=True, size=10); c.fill = HDR; c.border = BD; c.alignment = CT
            ws.column_dimensions[L(j + 1)].width = w
        for i, r in enumerate(rows):
            mx = 1
            for j, v in enumerate(r):
                c = ws.cell(row=i + 2, column=j + 1, value=v); c.font = Font(name='Meiryo UI', size=10); c.border = BD; c.alignment = WR
                f = fills(i, j, v) if fills else None
                if f: c.fill = f
                mx = max(mx, len(str(v or '')) / max(cols[j][1] * 0.9, 1) + str(v or '').count('\n'))
            ws.row_dimensions[i + 2].height = min(15 * (int(mx) + 1) + 4, 300)
        ws.freeze_panes = 'B2'; ws.auto_filter.ref = 'A1:%s%d' % (L(len(cols)), len(rows) + 1)
    def st_fill(v):
        v = str(v)
        return OK if '一致' in v or v == '登録済' else (RED if '未登録' in v or 'ない' in v else (WARN if '違う' in v else None))

    ws = wb.active; ws.title = 'メッセージ'
    rows = [[m['id'], m['type'], m['where'], m.get('ja_c', ''), m.get('ja_o', ''), m['vi'], m.get('ml') or '（なし）', m.get('mlText', ''), m['mlState'],
             where_text(m['use']) or '（どの画面でもまだ使っていない）', '\n'.join(x for x in [m.get('diff', ''), (m.get('open') or ['', ''])[0]] if x)] for m in d['messages']]
    sheet(ws, [('ID', 6), ('種類', 7), ('出す場所', 18), ('法人・拠点向け（丁寧）', 42), ('社内向け', 42), ('Tiếng Việt', 42), ('Message List', 10),
               ('Message List の文言（自動で読んだ）', 40), ('状態', 18), ('使う画面・項目', 40), ('未確認・OPEN', 40)], rows,
          lambda i, j, v: st_fill(v) if j == 8 else (WARN if j == 10 and v else None))

    ws = wb.create_sheet('メール')
    rows = [[m['id'], m['ja'], m['vi'], m.get('ml') or '（なし）', m.get('mlSubject', ''), m['mlState'], m.get('to', ''), m.get('when', ''),
             where_text(m['use']) or '（どの画面でもまだ使っていない）', m.get('src', '')] for m in d['mails']]
    sheet(ws, [('ID', 6), ('メール', 34), ('Email (VN)', 34), ('Message List', 10), ('Message List の件名（自動で読んだ）', 40), ('状態', 14),
               ('宛先', 30), ('送るとき', 34), ('使う画面・項目', 34), ('文面の正', 30)], rows, lambda i, j, v: st_fill(v) if j == 5 else None)

    ws = wb.create_sheet('ML のメール（設計書で未使用）')
    rows = [['MAIL-%d' % m['no'], m['use'], m['to'], m['subject']] for m in d['mlOnly']]
    sheet(ws, [('Message List', 10), ('利用ケース', 50), ('送信対象', 46), ('件名', 50)], rows)

    ws = wb.create_sheet('Message List への追加案')
    rows = [['', a['phase'], a['scene'], a['screen'], a['button'], a['title'], a['ja'], a['popup'], a['type'], a['fmt'], a['id']] for a in d['adds']]
    sheet(ws, [('No', 5), ('フェーズ', 9), ('場面', 30), ('画面', 30), ('ボタン', 18), ('メッセージタイトル', 22), ('メッセージ（JP）', 50),
               ('ポップアップのボタン', 24), ('メッセージタイプ', 14), ('メッセージ表示形式', 18), ('（参考）設計書のID', 10)], rows)
    note = wb.create_sheet('この一覧について')
    note.column_dimensions['A'].width = 120
    for i, t in enumerate([
        'メッセージ・メールの一覧（すべての画面設計書をまたぐ）。make.py を実行するたびに作り直す。',
        '元データ：05_基本設計/データ/_共通_メッセージ.py（ID・文言・Message List の番号）。各画面のデータは ID だけ使う。',
        'Message List（正）：%s を自動で読み、登録済み（一致）／登録済み（文言が違う）／未登録 を判定した。' % (d['mlFile'] or '（見つからなかった）'),
        '「Message List への追加案」は Message List の「運営管理者／法人顧客向けWEB」シートと同じ列の順。No は Message List に貼るときに振る。',
        '対象の画面設計書：' + '、'.join(x['title'][0] for x in d['docs']),
        'どの画面でもまだ使っていない ID：' + ('、'.join(d['unused']) or 'なし')]):
        c = note.cell(row=i + 1, column=1, value=t); c.font = Font(name='Meiryo UI', size=11); c.alignment = WR
    wb.save(path)

def build_html(d, path):
    tpl = open(os.path.join(HERE, 'common_template.html'), encoding='utf-8').read()
    js = json.dumps(d, ensure_ascii=False).replace('</', '<\\/')
    open(path, 'w', encoding='utf-8').write(tpl.replace('/*__DATA__*/', 'window.KC=' + js + ';'))
