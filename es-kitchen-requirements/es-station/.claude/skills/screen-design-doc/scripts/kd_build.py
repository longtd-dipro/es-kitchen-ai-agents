# -*- coding: utf-8 -*-
"""
画面設計書の Excel（日本語版・ベトナム語版を別ファイル）と HTML を作る（make.py から呼ぶ）。
  - 1機能（データファイル1つ）＝1シート。画面・状態はシートの中に上から順に並べる（状態ごとにシートを分けない）
  - 共通パターン（一覧・検索・削除の確認など）は「共通パターン」シートに1回だけ書き、画面には差分だけ書く
  - 列「データ例」：開発（ベトナム）向けに、入る値の例と意味
"""
import json, os, re, datetime
import kd_rules as R

HERE = os.path.dirname(os.path.abspath(__file__))

KIND = {
    "area": ["エリア", "Khu vực"], "button": ["ボタン", "Nút"], "link": ["リンク", "Liên kết"],
    "text": ["テキスト入力", "Ô nhập văn bản"], "textarea": ["複数行入力", "Ô nhập nhiều dòng"],
    "select": ["プルダウン", "Dropdown"], "month": ["年月", "Năm-tháng"], "date": ["日付", "Ngày"],
    "check": ["チェックボックス", "Checkbox"], "radio": ["ラジオボタン", "Radio"],
    "table": ["表", "Bảng"], "label": ["表示", "Hiển thị"], "tab": ["タブ", "Tab"],
    "toast": ["トースト", "Toast"], "file": ["ファイル添付", "Đính kèm tệp"], "modal": ["モーダル", "Modal"],
    "card": ["選択カード", "Thẻ chọn"],
}
TRIG = {
    "click": ["クリック", "Nhấp"], "select": ["選択", "Chọn"], "input": ["入力", "Nhập"],
    "view": ["閲覧のみ", "Chỉ xem"], "check": ["チェック", "Tích chọn"], "show": ["表示時", "Khi hiển thị"],
    "upload": ["ファイル選択", "Chọn tệp"],
}

# 見出し（0＝日本語版、1＝ベトナム語版）
LB = {
    'doc': ['画面設計書（基本設計）', 'Thiết kế màn hình (Thiết kế cơ bản)'],
    'draft': ['ドラフト：未確認の点が残っています。お客様・開発には渡さないでください。', 'BẢN NHÁP: còn điểm chưa xác nhận. Không gửi cho khách hàng / dev.'],
    'cover': ['表紙', 'Bìa'], 'list': ['画面一覧', 'Danh sách màn hình'], 'pat': ['共通パターン', 'Mẫu chung'],
    'msg': ['メッセージ一覧', 'Danh sách thông báo'], 'mail': ['メール', 'Email'], 'open': ['OPEN（お客様確認待ち）', 'OPEN (chờ khách xác nhận)'],
    'log': ['確認ログ', 'Nhật ký xác nhận'], 'rev': ['改訂履歴', 'Lịch sử sửa đổi'],
    'system': ['システム', 'Hệ thống'], 'func': ['機能', 'Chức năng'], 'screens': ['対象の画面', 'Màn hình'],
    'created': ['作成日', 'Ngày tạo'], 'author': ['作成者', 'Người tạo'], 'demo': ['元にしたデモ', 'Demo gốc'],
    'howto': ['作り方', 'Cách tạo'],
    'howto_t': ['データ（%s）が正。make.py でデモに番号を重ねた画像と Excel・HTML を作る。直すときはデータを直して作り直す（Excel を直接直さない）。',
                'Dữ liệu (%s) là bản gốc. make.py tạo ảnh đánh số trên demo, Excel và HTML. Muốn sửa thì sửa dữ liệu rồi tạo lại (không sửa trực tiếp Excel).'],
    'color': ['番号の色', 'Màu số'], 'color_t': ['青＝エリア（大項目）、赤＝項目（小項目）', 'Xanh = khu vực (mục lớn), Đỏ = mục (mục nhỏ)'],
    'code': ['Screen Code', 'Screen Code'],
    'screen': ['画面', 'Màn hình'], 'state': ['状態', 'Trạng thái'], 'note': ['補足', 'Ghi chú'],
    'layout': ['画面レイアウト（デモ %s に番号を重ねた画像）', 'Bố cục màn hình (ảnh demo %s có đánh số)'],
    'layout_web': ['画面レイアウト（本物の Web %s に番号を重ねた画像）', 'Bố cục màn hình (ảnh web thật %s có đánh số)'],
    'items': ['項目定義', 'Định nghĩa mục'], 'nitems': ['項目数', 'Số mục'], 'nopen': ['OPEN', 'OPEN'], 'go': ['場所', 'Vị trí'],
    'pat_use': ['【共通パターン %s】', '【Mẫu chung %s】'],
    'ver': ['版', 'Phiên bản'], 'preview': ['下見（開発に未提出）', 'Bản xem trước (chưa giao dev)'], 'hist': ['改訂履歴', 'Lịch sử sửa đổi'], 'demo_ok': ['デモとの差（確認済み）：', 'Khác demo (đã xác nhận): '],
}
COLS = [  # キー, 見出し JA, 見出し VI, 幅
    ('no', 'No', 'No', 7), ('chg', '変更', 'Thay đổi', 9), ('name', '項目名', 'Tên mục', 24), ('kind', '種類', 'Loại', 11), ('req', '必須', 'Bắt buộc', 8),
    ('len', '型・桁数', 'Kiểu / độ dài', 16), ('init', '初期値', 'Mặc định', 16), ('ex', 'データ例', 'Dữ liệu ví dụ', 26),
    ('detail', '詳細', 'Chi tiết', 48), ('cond', '表示・活性の条件', 'Điều kiện hiển thị', 24), ('trig', 'トリガー', 'Trigger', 10),
    ('valid', 'バリデーション', 'Validation', 30), ('err', 'メッセージ', 'Thông báo', 40), ('open', 'OPEN', 'OPEN', 30),
    ('demo', 'デモの定義（自動取得）', 'Định nghĩa từ demo (tự động)', 30)]

def L(p, lang):
    if isinstance(p, (list, tuple)): return (p[lang] if len(p) > lang and p[lang] else p[0]) or ''
    return p or ''

# ---------------------------------------------------------------- モデル
def msg_text(D, code, name, ln):
    m = D.MESSAGES.get(code)
    if not m:
        mm = D.MAILS.get(code)
        return ['%s：%s（宛先：%s）' % (code, mm['ja'], mm['to']), '%s: %s (gửi: %s)' % (code, mm['vi'], mm.get('to_vi', mm['to']))] if mm else [code, code]
    n = re.search(r'([\d,]+)', ln or '')
    ja = m['ja'].replace('{項目名}', name[0]); vi = m['vi'].replace('{tên mục}', name[1])
    if m.get('n_from_len') and n:
        ja = ja.replace('{n}', n.group(1)); vi = vi.replace('{n}', n.group(1))
    return ['%s：%s' % (code, ja), '%s: %s' % (code, vi)]

def build_model(D, rep, P, img_dir):
    pats = getattr(P, 'PATTERNS', {}) if P else {}
    views, issues, used_pats = [], [], {}
    for v in D.VIEWS:
        rv = rep['views'].get(v['id'], {})
        rows = []
        for it in v['items']:
            name = [it['ja'], it['vi']]
            ln = R.len_pair(it.get('len', ''))
            demo = rep['fields'].get(v['id'] + ':' + it['no'])
            demo_txt, auto = ['', ''], []
            if demo:
                d = demo['demo']
                demo_txt = ['項目名：%s／必須表示：%s／部品：%s' % (d['label'], 'あり' if d['req'] == 'true' else 'なし', d['type']),
                            'Tên: %s / Hiện "必須": %s / Loại: %s' % (d['label'], 'có' if d['req'] == 'true' else 'không', d['type'])]
                if d['opts']:
                    o = '・'.join(x for x in d['opts'] if x); demo_txt = [demo_txt[0] + '／選択肢：' + o, demo_txt[1] + ' / Lựa chọn: ' + o]
                if d['label'] != it['ja']: auto.append(['項目名がデモ（%s）と違う' % d['label'], 'Tên khác demo (%s)' % d['label']])
                rq = it.get('req', '')
                if rq == '○' and d['req'] != 'true': auto.append(['仕様は必須だが、デモに「必須」の表示がない', 'Spec bắt buộc nhưng demo không hiện "必須"'])
                if rq == '－' and d['req'] == 'true': auto.append(['仕様は任意だが、デモは「必須」の表示', 'Spec tùy chọn nhưng demo hiện "必須"'])
                if re.search(r'文字列 \d', ln[0]) and not d.get('maxlength'):
                    auto.append(['デモに最大文字数の制御がない（仕様 %s）' % ln[0], 'Demo chưa giới hạn số ký tự (spec %s)' % ln[1]])
            found = it['no'] not in rv.get('missing', [])
            if not found: auto.append(['デモの画面で場所が見つからない', 'Không tìm thấy vị trí trên demo'])
            ok = R.pair(it.get('demo_ok'))
            if ok[0]:                                   # demo_ok は「場所が見つからない」（見本で作れない状態）も確認済みとして扱う
                demo_txt = [(demo_txt[k] + '\n' if demo_txt[k] else '') + LB['demo_ok'][k] + (ok[k] or ok[0]) for k in (0, 1)]
                auto = []
            detail = R.pair(it.get('detail'))
            if it.get('pattern'):
                pid = it['pattern']; used_pats[pid] = pats[pid]
                detail = [LB['pat_use'][0] % pid + detail[0], LB['pat_use'][1] % pid + detail[1]]
            rq = it.get('req', '－' if it['kind'] in R.INPUT_KINDS else '')
            row = {'no': it['no'], 'lvl': 0 if '.' not in it['no'] else 1, 'ja': it['ja'], 'vi': it['vi'], 'sel': it['sel'],
                   'kindKey': it['kind'], 'kind': KIND[it['kind']], 'req': rq, 'reqP': [rq, R.REQ_VI.get(rq, rq)],
                   'len': ln, 'init': R.pair(it.get('init')), 'cond': R.pair(it.get('cond')), 'trig': TRIG[it.get('trig', 'view')],
                   'valid': R.pair(it.get('valid')), 'detail': detail, 'ex': R.pair(it.get('ex')),
                   'err': [msg_text(D, c, name, ln[0]) for c in it.get('err', [])], 'errCodes': it.get('err', []),
                   'chk': R.pair(it.get('open')), 'auto': auto, 'demo': demo_txt, 'found': found}
            rows.append(row)
            if row['chk'][0]:
                issues.append({'view': v['id'], 'code': v['code'], 'no': it['no'], 'ja': it['ja'], 'vi': it['vi'],
                               'src': 'OPEN（%s）' % it.get('ask', 'お客様'), 'text': row['chk'], 'ask': it.get('ask', '')})
        scr = [s for s in D.SCREENS if s['code'] == v['code']][0]
        views.append({'id': v['id'], 'code': v['code'], 'screen': [scr['ja'], scr['vi']], 'state': v['state'],
                      'note': R.pair(v.get('note')), 'setup': v.get('setup', ''), 'hide': v.get('hide', []),
                      'url': v.get('url', ''), 'snapshot': rv.get('snapshot'),
                      'img': os.path.join(img_dir, 'img', '%s_%s.png' % (D.IMG_PREFIX, v['id'])),
                      'imgRel': 'img/%s_%s.png' % (D.IMG_PREFIX, v['id']), 'rows': rows, 'missing': rv.get('missing', [])})
    for code, m in D.MESSAGES.items():
        if m.get('open'):
            o = R.pair(m['open'])
            issues.append({'view': '', 'code': 'メッセージ', 'no': code, 'ja': m['ja'][:30], 'vi': m['vi'][:30], 'src': 'OPEN（文言）', 'text': o, 'ask': ''})
    return {'views': views, 'issues': issues, 'patterns': used_pats}

# ---------------------------------------------------------------- 書き出し
def write(built, book, only, draft, ver=None, hist=None):
    files = []
    groups = [(book, built)] if book else [(None, [b]) for b in built]
    for name, items in groups:
        out = items[0][3]
        base = name or items[0][0].BASENAME
        if only != 'html':
            for lang, suf in ((0, 'JA'), (1, 'VI')):
                p = os.path.join(out, '%s_%s.xlsx' % (base, suf)); build_excel(p, items, lang, draft, name, ver, hist or []); files.append(p)
        if only != 'excel':
            for D, model, demo, o in items:
                p = os.path.join(o, (D.BASENAME + ('_ドラフト' if draft else '')) + '.html'); build_html(p, D, model, demo, draft); files.append(p)
    return files

def sheet_title(D, lang, used):
    t = (D.SHEET[lang] if isinstance(D.SHEET, (list, tuple)) else D.SHEET) if D.SHEET else L(D.TITLE, lang)
    t = re.sub(r'[\[\]\*\?/\\:：]', ' ', t)[:31]
    s, i = t, 2
    while s in used: s = t[:28] + '(%d)' % i; i += 1
    used.add(s); return s

def build_excel(path, items, lang, draft, book, ver=None, hist=()):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XImg
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as CL
    wb = Workbook()
    HDR = PatternFill('solid', fgColor='92CDDC'); SEC = PatternFill('solid', fgColor='1F4E79'); AREA = PatternFill('solid', fgColor='EAF1FB')
    WARN = PatternFill('solid', fgColor='FFF2CC'); RED = PatternFill('solid', fgColor='FCE4E4')
    thin = Side(style='thin', color='A6B4C3'); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
    WR = Alignment(wrap_text=True, vertical='top'); CT = Alignment(horizontal='center', vertical='top', wrap_text=True)
    NW = Alignment(vertical='center', wrap_text=False)
    F = 'Meiryo UI' if lang == 0 else 'Arial'
    def cell(ws, r, c, v, bold=False, fill=None, al=WR, size=10, border=True, color=None):
        x = ws.cell(row=r, column=c, value=v)
        x.font = Font(name=F, bold=bold, size=size, color=color); x.alignment = al
        if fill: x.fill = fill
        if border: x.border = BD
        return x
    def table(ws, heads, rows, r0=1, fills=None):
        for j, (h, w) in enumerate(heads):
            cell(ws, r0, j + 1, h, True, HDR, CT); ws.column_dimensions[CL(j + 1)].width = w
        for i, row in enumerate(rows):
            mx = 1
            for j, v in enumerate(row):
                c = cell(ws, r0 + 1 + i, j + 1, v)
                f = fills(i, j, v) if fills else None
                if f: c.fill = f
                mx = max(mx, len(str(v or '')) / max(heads[j][1] * 0.9, 1) + str(v or '').count('\n'))
            ws.row_dimensions[r0 + 1 + i].height = min(15 * (int(mx) + 1) + 4, 300)
        ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    D0 = items[0][0]
    used = set()
    titles = [sheet_title(D, lang, used) for D, _, _, _ in items]

    # 表紙
    ws = wb.active; ws.title = LB['cover'][lang]
    ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 22; ws.column_dimensions['C'].width = 90
    cell(ws, 2, 2, LB['doc'][lang], True, border=False, size=18)
    if draft: cell(ws, 3, 2, LB['draft'][lang], True, border=False, size=12, color='C00000', al=NW)
    rows = [(LB['ver'][lang], ver or LB['preview'][lang]), (LB['system'][lang], L([D0.SYSTEM['ja'], D0.SYSTEM['vi']], lang)),
            (LB['func'][lang], book or L(D0.TITLE, lang)),
            (LB['screens'][lang], '\n'.join('%s %s' % (s['code'], L([s['ja'], s['vi']], lang)) for D, *_ in items for s in D.SCREENS)),
            (LB['created'][lang], D0.CREATED), (LB['author'][lang], D0.AUTHOR),
            (LB['demo'][lang], '\n'.join(D.DEMO_FILE for D, *_ in items)),
            (LB['howto'][lang], LB['howto_t'][lang] % '、'.join(os.path.basename(D.__file__) for D, *_ in items)),
            (LB['color'][lang], LB['color_t'][lang]), (LB['code'][lang], L(getattr(D0, 'CODE_NOTE', ''), lang))]
    for i, (k, v) in enumerate(rows):
        cell(ws, 5 + i, 2, k, True, HDR); c = cell(ws, 5 + i, 3, v)
        ws.row_dimensions[5 + i].height = max(18, 15 * (str(v).count('\n') + 1) + 4)
    r0 = 5 + len(rows) + 2
    cell(ws, r0, 2, LB['rev'][lang], True, border=False, size=12)
    for j, h in enumerate([['版', 'Phiên bản'], ['日付', 'Ngày'], ['内容', 'Nội dung'], ['作成者', 'Người tạo']]):
        cell(ws, r0 + 1, 2 + j, h[lang], True, HDR)
    ws.column_dimensions['D'].width = 50; ws.column_dimensions['E'].width = 20
    revs = [[h['ver'], h.get('date') or datetime.date.today().strftime('%Y-%m-%d'), '、'.join(h.get('files', [])), D0.AUTHOR] for h in reversed(list(hist))]
    if not revs: revs = [[L(x, lang) for x in rv] for D, *_ in items for rv in (D.REVISIONS or [])]
    for i, rv in enumerate(revs):
        for j, v in enumerate(rv): cell(ws, r0 + 2 + i, 2 + j, v)

    # 改訂履歴（開発が変更に気づけるように。新しい版が上）
    TYPE = {'追加': ['追加', 'Thêm'], '変更': ['変更', 'Sửa'], '削除': ['削除', 'Xóa']}
    ws = wb.create_sheet(LB['hist'][lang])
    chrows = []
    for D, *_ in items:
        for c in D.CHANGES:
            nos = c['no'] if isinstance(c['no'], list) else [c['no']]
            chrows.append([c['ver'], c.get('date', ''), L(D.TITLE, lang), '、'.join('全体' if n == '*' else n for n in nos),
                           L(TYPE.get(c.get('type', '変更'), [c.get('type', '')] * 2), lang),
                           L(c.get('before'), lang), L(c.get('after'), lang), L(c.get('reason'), lang), L(c.get('impact'), lang)])
    chrows.sort(key=lambda r: [int(x) for x in re.findall(r'\d+', r[0])], reverse=True)
    if hist: chrows.append([hist[0]['ver'], hist[0].get('date', ''), '、'.join(hist[0].get('files', [])), '全体', L(['初版', 'Bản đầu'], lang), '', '', '', ''])
    table(ws, [(LB['ver'][lang], 7), (['日付', 'Ngày'][lang], 11), (LB['func'][lang], 24), (['項目No', 'Mục'][lang], 12), (['種類', 'Loại'][lang], 7),
               (['変更前', 'Trước'][lang], 34), (['変更後', 'Sau'][lang], 34), (['理由', 'Lý do'][lang], 30), (['開発への影響', 'Ảnh hưởng tới dev'][lang], 34)],
          chrows, fills=lambda i, j, v: WARN if ver and chrows[i][0] == ver else None)

    # 画面一覧（リンクは機能シートの状態の位置へ）
    ws_list = wb.create_sheet(LB['list'][lang])
    heads = [('No', 5), ('Screen Code', 16), (LB['screen'][lang], 36), (LB['state'][lang], 40), (LB['func'][lang], 28),
             (LB['nitems'][lang], 8), (LB['nopen'][lang], 8), (LB['note'][lang], 60)]
    for j, (h, w) in enumerate(heads):
        cell(ws_list, 1, j + 1, h, True, HDR, CT); ws_list.column_dimensions[CL(j + 1)].width = w
    ws_list.freeze_panes = 'A2'
    list_row = 2

    # 共通パターン
    pats = {}
    for _, m, _, _ in items: pats.update(m['patterns'])
    if pats:
        ws = wb.create_sheet(LB['pat'][lang])
        rows = []
        for pid, p in pats.items():
            rows.append([pid, L(p['name'], lang), '\n'.join('・' + L(x, lang) for x in p.get('rules', [])), ', '.join(p.get('msgs', []))])
        table(ws, [('ID', 10), (LB['pat'][lang], 26), (['決まり', 'Quy tắc'][lang], 90), (['メッセージ', 'Thông báo'][lang], 20)], rows)

    # 追加のシート（初期データなど。データの SHEETS＝[{name:[ja,vi], note:[ja,vi], tables:[{title:[ja,vi], head:[[[ja,vi], 幅], …], rows:[[値 or [ja,vi], …], …]}]}]）
    for D, *_ in items:
        for sh in getattr(D, 'SHEETS', []) or []:
            ws = wb.create_sheet(re.sub(r'[\[\]\*\?/\\:：]', ' ', L(sh['name'], lang))[:31])
            r = 1
            if sh.get('note'):
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
                cell(ws, r, 1, L(sh['note'], lang), border=False, color='44546A'); ws.row_dimensions[r].height = 15 * (len(L(sh['note'], lang)) // 110 + 2); r += 2
            for tb in sh.get('tables', []):
                cell(ws, r, 1, L(tb['title'], lang), True, border=False, size=12); r += 1
                if tb.get('note') and L(tb['note'], lang):
                    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
                    cell(ws, r, 1, L(tb['note'], lang), border=False, color='44546A'); ws.row_dimensions[r].height = 15 * (len(L(tb['note'], lang)) // 110 + 2); r += 1
                table(ws, [(L(h[0], lang), h[1]) for h in tb['head']], [[L(x, lang) if isinstance(x, (list, tuple)) else x for x in row] for row in tb['rows']], r0=r)
                r += len(tb['rows']) + 3
            ws.freeze_panes = None

    # 機能シート（状態を上から並べる）
    IMGW = 760; IMGC = 2; NIMGC = 13; C0 = IMGC + NIMGC + 1
    for (D, model, demo, out), title in zip(items, titles):
        ws = wb.create_sheet(title); ws.sheet_view.zoomScale = 85
        for c in range(1, C0): ws.column_dimensions[CL(c)].width = 9.0
        ws.column_dimensions['A'].width = 2
        for j, col in enumerate(COLS): ws.column_dimensions[CL(C0 + j)].width = col[3]
        last = C0 + len(COLS) - 1
        cell(ws, 1, 2, '%s ｜ %s ｜ %s' % (LB['doc'][lang], L([D.SYSTEM['ja'], D.SYSTEM['vi']], lang), L(D.TITLE, lang)), True, border=False, size=14, al=NW)
        if draft: cell(ws, 2, 2, LB['draft'][lang], True, border=False, color='C00000', al=NW)
        r = 4
        for v in model['views']:
            # 見出し帯
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last)
            cell(ws, r, 2, '■ %s %s ｜ %s %s' % (v['code'], L(v['screen'], lang), v['id'], L(v['state'], lang)), True, SEC, NW, 11, color='FFFFFF')
            ws.row_dimensions[r].height = 22
            nopen = sum(1 for x in model['issues'] if x['view'] == v['id'])
            for j, x in enumerate([list_row - 1, v['code'], L(v['screen'], lang), '%s %s' % (v['id'], L(v['state'], lang)), title, len(v['rows']), nopen, L(v['note'], lang)]):
                c = cell(ws_list, list_row, j + 1, x, al=CT if j in (0, 5, 6) else WR)
                if j == 3: c.hyperlink = "#'%s'!B%d" % (title, r); c.font = Font(name=F, color='1D4ED8', underline='single', size=10)
            list_row += 1
            r += 1
            if L(v['note'], lang):
                ws.merge_cells(start_row=r, start_column=C0, end_row=r, end_column=last)
                cell(ws, r, C0, LB['note'][lang] + '：' + L(v['note'], lang), border=False, color='44546A')
                ws.row_dimensions[r].height = max(18, 15 * (len(L(v['note'], lang)) // 150 + 1) + 4)
                r += 1
            if R.is_url(D.DEMO_FILE): cell(ws, r, IMGC, LB['layout_web'][lang] % v['url'], True, border=False, al=NW)
            else: cell(ws, r, IMGC, LB['layout'][lang] % os.path.basename(D.DEMO_FILE), True, border=False, al=NW)
            for j, col in enumerate(COLS): cell(ws, r, C0 + j, col[1 + lang], True, HDR, CT)
            r += 1; r_img = r; px = 0
            cur = [c for c in D.CHANGES if ver and c.get('ver') == ver]
            chg = {}
            for c in cur:
                for n in (c['no'] if isinstance(c['no'], list) else [c['no']]): chg[n] = c.get('type', '変更')
            for row in v['rows']:
                area = row['lvl'] == 0
                t = chg.get(row['no']) or chg.get('*')
                chg_t = ('%s %s' % (ver, {'追加': ['追加', 'Thêm'], '変更': ['変更', 'Sửa'], '削除': ['削除', 'Xóa']}.get(t, [t, t])[lang])) if t else ''
                ex = row['ex'][0] if lang == 0 else ' → '.join(x for x in row['ex'] if x)
                open_t = L(row['chk'], lang)
                if row['auto']: open_t = (open_t + '\n' if open_t else '') + '\n'.join('[!] ' + L(a, lang) for a in row['auto'])
                vals = {'no': row['no'], 'chg': chg_t, 'name': ('' if area else '　') + (row['ja'] if lang == 0 else row['vi']), 'kind': L(row['kind'], lang),
                        'req': L(row['reqP'], lang), 'len': L(row['len'], lang), 'init': L(row['init'], lang), 'ex': ex,
                        'detail': L(row['detail'], lang), 'cond': L(row['cond'], lang), 'trig': L(row['trig'], lang),
                        'valid': L(row['valid'], lang), 'err': '\n'.join(L(e, lang) for e in row['err']), 'open': open_t, 'demo': L(row['demo'], lang)}
                mx = 1
                for j, col in enumerate(COLS):
                    x = vals[col[0]]
                    c = cell(ws, r, C0 + j, x or ('－' if col[0] in ('req', 'len', 'init', 'cond', 'valid', 'err') else ''),
                             bold=area and j < 2, fill=AREA if area else None, al=CT if col[0] in ('no', 'chg', 'kind', 'req', 'trig') else WR)
                    if col[0] == 'open' and x: c.fill = WARN
                    if chg_t and col[0] in ('no', 'chg', 'name'): c.fill = PatternFill('solid', fgColor='FFD8A8')
                    if col[0] == 'ex' and x and not area: c.fill = PatternFill('solid', fgColor='EEF7EE')
                    mx = max(mx, len(str(x or '')) / max(col[3] * 0.55, 1) + str(x or '').count('\n'))
                h = min(15 * (int(mx) + 1) + 4, 300); ws.row_dimensions[r].height = h; px += h * 96 / 72
                r += 1
            # 画像（セルに固定。Google スプレッドシートでも出る）
            if os.path.exists(v['img']):
                img = XImg(v['img']); sc = min(1.0, IMGW / img.width); img.width = int(img.width * sc); img.height = int(img.height * sc)   # 狭い画像（スマホのアプリ）は拡大しない
                while px < img.height:
                    ws.row_dimensions[r].height = 15; px += 20; r += 1
                anchor_img(ws, img, IMGC, r_img, img.width, img.height)
            r += 2

    # メッセージ一覧
    msgs = {}
    for D, *_ in items: msgs.update(D.MESSAGES)
    ws = wb.create_sheet(LB['msg'][lang])
    rows = [[k, m['type'], m['where'], m['ja'] if lang == 0 else m['vi'], m['vi'] if lang == 0 else m['ja'], m.get('ml') or '－', L(m.get('open'), lang)]
            for k, m in sorted(msgs.items())]
    table(ws, [('ID', 7), (['種類', 'Loại'][lang], 9), (['出す場所', 'Vị trí'][lang], 20), (['文言', 'Nội dung'][lang], 56),
               (['文言（VN）', 'Nội dung (JP)'][lang], 56), ('Message List', 12), ('OPEN', 30)], rows,
          fills=lambda i, j, v: WARN if j == 6 and v else None)

    mails = {}
    for D, *_ in items: mails.update(D.MAILS)
    if mails:
        ws = wb.create_sheet(LB['mail'][lang])
        rows = [[k, L([m['ja'], m['vi']], lang), m.get('when', ''), m['to'], m.get('ml', ''), m.get('src', '')] for k, m in mails.items()]
        table(ws, [('ID', 7), (['メール', 'Email'][lang], 44), (['送るとき', 'Khi gửi'][lang], 36), (['宛先', 'Người nhận'][lang], 30),
                   ('Message List', 12), (['文面の正', 'Nội dung gốc'][lang], 36)], rows)

    ws = wb.create_sheet(LB['open'][lang])
    iss = [x for _, m, _, _ in items for x in m['issues']]
    rows = [[i + 1, x['code'], x['view'], x['no'], x['ja'] if lang == 0 else x['vi'], x.get('ask', ''), L(x['text'], lang), ''] for i, x in enumerate(iss)]
    table(ws, [('No', 5), (LB['screen'][lang], 16), (LB['state'][lang], 9), (['項目No', 'Mục'][lang], 8), (['項目名', 'Tên mục'][lang], 26),
               (['確認先', 'Hỏi ai'][lang], 14), (['内容', 'Nội dung'][lang], 70), (['回答', 'Trả lời'][lang], 40)], rows)

    ws = wb.create_sheet(LB['log'][lang])
    decs = [d for D, *_ in items for d in D.DECISIONS]
    rows = [[i + 1, d['date'], d['target'], L(d['q'], lang), L(d['a'], lang), d['src']] for i, d in enumerate(decs)]
    table(ws, [('No', 5), (['日付', 'Ngày'][lang], 12), (['対象', 'Đối tượng'][lang], 22), (['確認したこと', 'Câu hỏi'][lang], 60),
               (['決まったこと', 'Kết luận'][lang], 60), (['根拠', 'Căn cứ'][lang], 30)], rows)
    wb.save(path)

def anchor_img(ws, img, col0, row0, wpx, hpx):
    from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
    from openpyxl.utils.units import pixels_to_EMU
    from openpyxl.utils import get_column_letter as CL
    c, rem = col0, wpx
    while True:
        px = int((ws.column_dimensions[CL(c)].width or 8.43) * 7 + 5)
        if rem <= px: break
        rem -= px; c += 1
    ec, ecoff = c, rem
    r, rem = row0, hpx
    while True:
        px = (ws.row_dimensions[r].height or 15) * 96 / 72
        if rem <= px: break
        rem -= px; r += 1
    img.anchor = TwoCellAnchor(editAs='oneCell',
        _from=AnchorMarker(col=col0 - 1, colOff=0, row=row0 - 1, rowOff=0),
        to=AnchorMarker(col=ec - 1, colOff=pixels_to_EMU(int(ecoff)), row=r - 1, rowOff=pixels_to_EMU(int(rem))))
    ws.add_image(img)

# ---------------------------------------------------------------- HTML（デモを中に入れた1ファイル。日越切替）
STYLE_RE = re.compile(r'<style[^>]*>.*?</style>', re.S)

def dedupe_styles(snaps):
    """状態ごとのスナップショットに同じ <style>（Next.js の CSS など）が何度も入るので、2つ以上に出るものは1つにまとめ、
    スナップショットには <!--KDCSS:n--> を置く（viewer_template.html が表示のときに元に戻す）。HTML のほとんどはこの重複"""
    count = {}
    for h in snaps.values():
        for b in set(STYLE_RE.findall(h)): count[b] = count.get(b, 0) + 1
    shared = [b for b, n in count.items() if n > 1 and len(b) > 200]
    idx = {b: i for i, b in enumerate(shared)}
    out = {k: STYLE_RE.sub(lambda m: '<!--KDCSS:%d-->' % idx[m.group(0)] if m.group(0) in idx else m.group(0), h) for k, h in snaps.items()}
    return out, shared

def build_html(path, D, model, demo, draft):
    web = R.is_url(D.DEMO_FILE)
    # 本物の Web のときは、状態ごとに撮ったスナップショット（静止した HTML）を入れる（デモは1ファイルを入れて setup で状態を作る）
    demo_html = '' if web else open(demo, encoding='utf-8').read()
    snaps = {v['id']: open(v['snapshot'], encoding='utf-8').read() for v in model['views'] if web and v.get('snapshot') and os.path.exists(v['snapshot'])}
    snaps, shared_css = dedupe_styles(snaps)
    overlay = open(os.path.join(HERE, 'overlay.js'), encoding='utf-8').read()
    tpl = open(os.path.join(HERE, 'viewer_template.html'), encoding='utf-8').read()
    views = [dict({k: x for k, x in v.items() if k not in ('img', 'snapshot')}, img=v['imgRel']) for v in model['views']]
    data = {'title': D.TITLE[0] + ('（ドラフト）' if draft else ''), 'titleVi': D.TITLE[1] + (' (nháp)' if draft else ''),
            'dataFile': os.path.basename(D.__file__), 'system': D.SYSTEM, 'created': D.CREATED, 'author': D.AUTHOR, 'demo': D.DEMO_FILE,
            'viewerCss': getattr(D, 'VIEWER_CSS', '.gnav{position:static!important}'), 'screens': D.SCREENS, 'views': views,
            'vw': getattr(D, 'VIEWPORT', (1440, 900))[0],
            'issues': model['issues'], 'hideAlways': D.HIDE_ALWAYS,
            'messages': [dict(id=k, **v) for k, v in D.MESSAGES.items()], 'mails': [dict(id=k, **v) for k, v in D.MAILS.items()],
            'sheets': getattr(D, 'SHEETS', []) or []}
    js = lambda o: json.dumps(o, ensure_ascii=False).replace('</', '<\\/')
    html = (tpl.replace('/*__DATA__*/', 'window.KD_DATA=' + js(data) + ';')
               .replace('/*__OVERLAY__*/', 'window.KD_OVERLAY=' + js(overlay) + ';')
               .replace('/*__DEMO__*/', ('window.KD_CSS=' + js(shared_css) + ';window.KD_DEMOS=' + js(snaps) + ';') if web else ('window.KD_DEMO=' + js(demo_html) + ';')))
    open(path, 'w', encoding='utf-8').write(html)
